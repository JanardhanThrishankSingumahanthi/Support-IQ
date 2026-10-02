from __future__ import annotations

import mimetypes
import os
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, Depends, File, Form, HTTPException, Path as FastPath, Query, UploadFile, status
from fastapi.responses import FileResponse, Response
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db, get_pagination
from app.api.schemas import ErrorResponse, PaginatedResponse
from app.db.models import Document, User

router = APIRouter(prefix="/documents", tags=["documents"])

ALLOWED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".csv": "csv"}
MAX_DOCUMENT_SIZE_BYTES = 10 * 1024 * 1024


def document_storage_root() -> Path:
    from app.core.config import get_settings
    return get_settings().document_storage_dir


def normalize_document_status(value: str | None) -> str:
    status_value = (value or "draft").upper()
    valid = {"UPLOADING", "EXTRACTING", "CHUNKING", "EMBEDDING", "INDEXING", "COMPLETED", "FAILED", "DRAFT"}
    return status_value if status_value in valid else "DRAFT"


def serialize_document(document: Document) -> dict:
    metadata = document.metadata_json or {}
    indexed_at = metadata.get("indexed_at")
    return {
        "id": document.id,
        "title": document.title,
        "filename": metadata.get("filename") or document.title,
        "file_type": metadata.get("file_type") or "text",
        "size": int(metadata.get("size") or 0),
        "category": metadata.get("category") or "General",
        "version": int(metadata.get("version") or 1),
        "status": normalize_document_status(document.status),
        "owner_id": document.owner_id,
        "uploaded_at": document.created_at.isoformat(),
        "updated_at": document.updated_at.isoformat(),
        "indexed_at": indexed_at.isoformat() if isinstance(indexed_at, datetime) else indexed_at,
        "chunk_count": int(metadata.get("chunk_count") or 0),
        "preview": (document.content or "")[:400],
        "storage_path": metadata.get("storage_path"),
        "error": metadata.get("error"),
    }


def build_chunks(text: str) -> list[str]:
    if not text.strip():
        return []
    paragraphs = [part.strip() for part in text.splitlines() if part.strip()]
    if not paragraphs:
        return [text.strip()]
    chunks: list[str] = []
    current: list[str] = []
    for paragraph in paragraphs:
        current.append(paragraph)
        if len("\n".join(current)) >= 1000:
            chunks.append("\n".join(current).strip())
            current = []
    if current:
        chunks.append("\n".join(current).strip())
    return [chunk for chunk in chunks if chunk]


import io
import re
import xml.etree.ElementTree as ET
import zipfile
import zlib
from app.db.models import DocumentChunk


async def extract_text_bytes(file_name: str, payload: bytes) -> tuple[str, list[dict[str, Any]]]:
    """Extract text from supported file types, returning (full_text, page_records)."""
    extension = Path(file_name).suffix.lower()

    if extension in {".txt", ".csv"}:
        try:
            text = payload.decode("utf-8")
        except UnicodeDecodeError:
            text = payload.decode("latin1", errors="replace")
        return text, [{"page": 1, "text": text}]

    if extension == ".docx":
        # 1. Primary DOCX parser: python-docx
        try:
            import docx
            doc = docx.Document(io.BytesIO(payload))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            if paragraphs:
                full_text = "\n\n".join(paragraphs)
                return full_text, [{"page": 1, "text": full_text}]
        except Exception:
            pass

        # 2. Fallback XML parsing
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as z:
                if "word/document.xml" in z.namelist():
                    tree = ET.fromstring(z.read("word/document.xml"))
                    texts = [node.text for node in tree.iter() if node.text]
                    extracted = " ".join(texts).strip()
                    if extracted:
                        return extracted, [{"page": 1, "text": extracted}]
        except Exception:
            pass
        fallback_text = payload.decode("utf-8", errors="replace")
        return fallback_text, [{"page": 1, "text": fallback_text}]

    if extension == ".pdf":
        # 1. Primary PDF parser: pypdf (page-aware, handles compressed object streams)
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(payload))
            page_records: list[dict[str, Any]] = []
            for page_idx, page in enumerate(reader.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    page_records.append({"page": page_idx, "text": page_text})
            if page_records:
                full_text = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in page_records)
                return full_text, page_records
        except Exception:
            pass

        # 2. Secondary fallback: Stream decompression
        text_parts = []
        stream_pattern = re.compile(b"stream[\r\n]+(.*?)[\r\n]+endstream", re.DOTALL)
        for match in stream_pattern.finditer(payload):
            raw_stream = match.group(1)
            try:
                decompressed = zlib.decompress(raw_stream)
            except Exception:
                decompressed = raw_stream

            matches = re.findall(rb"\((.*?)\)\s*Tj", decompressed)
            if matches:
                decoded = " ".join([m.decode("latin1", errors="ignore") for m in matches if m.strip()])
                if decoded.strip():
                    text_parts.append(decoded.strip())
            else:
                bt_matches = re.findall(rb"\[(.*?)\]\s*TJ", decompressed)
                for bt in bt_matches:
                    inner = re.findall(rb"\((.*?)\)", bt)
                    if inner:
                        decoded = "".join([m.decode("latin1", errors="ignore") for m in inner if m.strip()])
                        if decoded.strip():
                            text_parts.append(decoded.strip())

        if text_parts:
            combined = "\n\n".join(text_parts)
            return combined, [{"page": 1, "text": combined}]

        return "", []

    if extension == ".docx":
        return "", []

    fallback = payload.decode("utf-8", errors="replace")
    return fallback, [{"page": 1, "text": fallback}]


async def run_document_pipeline(document: Document, payload: bytes, original_name: str, db: Session | None = None) -> Document:
    from app.retrieval.service import _hash_vector, tokenize

    stored_dir = document_storage_root()
    safe_name = os.path.basename(original_name)
    stored_path = stored_dir / f"{document.id}-{uuid4().hex}-{safe_name}"
    stored_path.write_bytes(payload)

    document.status = "EXTRACTING"
    document.metadata_json = {
        **(document.metadata_json or {}),
        "filename": safe_name,
        "storage_path": str(stored_path),
        "size": len(payload),
    }

    full_text, page_records = await extract_text_bytes(original_name, payload)
    if not full_text.strip() and len(payload) > 0:
        raise ValueError(f"No readable text could be extracted from '{safe_name}'. The file may be corrupt or image-only.")

    document.content = full_text
    document.metadata_json["file_type"] = Path(original_name).suffix.lower().lstrip(".") or "text"
    document.metadata_json["category"] = document.metadata_json.get("category") or "General"
    document.metadata_json["version"] = int(document.metadata_json.get("version") or 1)
    document.status = "CHUNKING"

    # Build page-aware chunks
    chunk_records: list[dict[str, Any]] = []
    for p_info in page_records:
        p_num = p_info.get("page", 1)
        p_text = p_info.get("text", "").strip()
        if not p_text:
            continue
        p_chunks = build_chunks(p_text)
        if not p_chunks:
            p_chunks = [p_text]
        for c_text in p_chunks:
            chunk_records.append({"page": p_num, "content": c_text})

    if not chunk_records and full_text.strip():
        chunk_records = [{"page": 1, "content": full_text.strip()}]

    document.metadata_json["chunk_count"] = len(chunk_records)
    document.status = "EMBEDDING"

    # Persist DocumentChunk records with pre-computed hash embeddings
    if db is not None:
        db.query(DocumentChunk).filter(DocumentChunk.document_id == document.id).delete()
        for idx, item in enumerate(chunk_records, start=1):
            c_text = item["content"]
            vector = _hash_vector(tokenize(c_text))
            db.add(
                DocumentChunk(
                    document_id=document.id,
                    version_id=None,
                    chunk_index=idx,
                    content=c_text,
                    metadata_json={
                        "page": item["page"],
                        "section": f"Section {idx}",
                        "tokens": len(c_text.split()),
                        "embedding": vector,
                    },
                )
            )

    document.status = "INDEXING"
    document.status = "COMPLETED"
    document.metadata_json["indexed_at"] = datetime.now(timezone.utc).isoformat()
    document.metadata_json["error"] = None
    return document



@router.get("", response_model=PaginatedResponse[dict])
def list_documents(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    pagination: dict = Depends(get_pagination),
    category: str | None = Query(default=None),
    status_filter: str | None = Query(default=None, alias="status"),
):
    query = db.query(Document)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name in ["Administrator", "Support Agent", "Agent"]))
    if not is_privileged:
        query = query.filter((Document.owner_id == current_user.id) | (Document.owner_id.is_(None)))
    if category:
        query = query.filter(Document.metadata_json.op("->>")("category") == category)
    if status_filter:
        query = query.filter(Document.status == status_filter)

    total = query.count()
    items = query.order_by(Document.updated_at.desc()).offset((pagination["page"] - 1) * pagination["page_size"]).limit(pagination["page_size"]).all()
    return {
        "items": [serialize_document(item) for item in items],
        "meta": {
            "page": pagination["page"],
            "page_size": pagination["page_size"],
            "total": total,
            "has_next": (pagination["page"] * pagination["page_size"]) < total,
            "has_previous": pagination["page"] > 1,
        },
    }


@router.get("/{document_id}")
def get_document(
    document_id: int = FastPath(..., gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Document).filter(Document.id == document_id)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name in ["Administrator", "Support Agent", "Agent"]))
    if not is_privileged:
        query = query.filter((Document.owner_id == current_user.id) | (Document.owner_id.is_(None)))
    document = query.first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "not_found", "message": "Document was not found."})
    doc_dict = serialize_document(document)
    doc_dict["content"] = document.content or ""
    doc_dict["chunks"] = [
        {
            "id": c.id,
            "chunk_index": c.chunk_index,
            "content": c.content,
            "page": (c.metadata_json or {}).get("page", c.chunk_index + 1),
            "metadata": c.metadata_json or {},
        }
        for c in sorted(document.chunks, key=lambda x: x.chunk_index)
    ]
    return doc_dict


@router.get("/{document_id}/download")
def download_document(
    document_id: int = FastPath(..., gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Document).filter(Document.id == document_id)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name in ["Administrator", "Support Agent", "Agent"]))
    if not is_privileged:
        query = query.filter((Document.owner_id == current_user.id) | (Document.owner_id.is_(None)))
    document = query.first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "not_found", "message": "Document was not found."})

    metadata = document.metadata_json or {}
    filename = metadata.get("filename") or f"{document.title}.txt"
    safe_filename = os.path.basename(filename)

    storage_path_str = metadata.get("storage_path")
    if storage_path_str:
        file_path = Path(storage_path_str).resolve()
        storage_root = document_storage_root().resolve()
        try:
            file_path.relative_to(storage_root)
        except ValueError:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail={"status": "access_denied", "message": "Invalid storage path."})

        if file_path.is_file():
            content_type, _ = mimetypes.guess_type(safe_filename)
            content_type = content_type or "application/octet-stream"
            return FileResponse(
                path=str(file_path),
                filename=safe_filename,
                media_type=content_type,
            )

    if document.content is not None:
        content_bytes = document.content.encode("utf-8")
        if not safe_filename.endswith(".txt") and not any(safe_filename.endswith(ext) for ext in ALLOWED_EXTENSIONS):
            safe_filename += ".txt"
        return Response(
            content=content_bytes,
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{safe_filename}"'},
        )

    raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "file_not_found", "message": "Original document file is not available."})


@router.post("", status_code=status.HTTP_201_CREATED)
async def upload_document(
    file: UploadFile = File(...),
    category: str = Form(default="General"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if file.filename is None or not file.filename.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail={"status": "invalid_file", "message": "A file name is required for upload."})

    extension = Path(file.filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={"status": "unsupported_file", "message": "Only PDF, DOCX, TXT, and CSV files are supported."},
        )

    payload = await file.read()
    if len(payload) == 0:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail={"status": "empty_file", "message": "Uploaded file is empty."})
    if len(payload) > MAX_DOCUMENT_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail={"status": "file_too_large", "message": "Maximum supported file size is 10 MB."},
        )

    safe_name = os.path.basename(file.filename or "")
    doc_title = Path(safe_name).stem or safe_name
    existing_doc = (
        db.query(Document)
        .filter(
            Document.owner_id == current_user.id,
            Document.title == doc_title,
            Document.status == "COMPLETED",
        )
        .first()
    )
    if existing_doc:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={"status": "duplicate", "message": f"A document named '{safe_name}' already exists in your knowledge base."},
        )

    document = Document(
        title=doc_title,
        content="",
        status="UPLOADING",
        owner_id=current_user.id,
        metadata_json={
            "filename": file.filename,
            "file_type": ALLOWED_EXTENSIONS[extension],
            "size": len(payload),
            "category": category.strip() or "General",
            "version": 1,
            "chunk_count": 0,
            "storage_path": None,
        },
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    try:
        document = await run_document_pipeline(document, payload, file.filename, db=db)
        document.metadata_json = {
            **(document.metadata_json or {}),
            "category": category.strip() or document.metadata_json.get("category") or "General",
        }
        document.status = "COMPLETED"
        db.add(document)
        db.commit()
        db.refresh(document)
        return serialize_document(document)
    except Exception as exc:  # pragma: no cover - preserves explicit failure state for processing errors.
        db.rollback()
        document.status = "FAILED"
        document.metadata_json = {
            **(document.metadata_json or {}),
            "category": category.strip() or "General",
            "error": str(exc),
        }
        db.add(document)
        db.commit()
        db.refresh(document)
        status_code = status.HTTP_400_BAD_REQUEST if isinstance(exc, ValueError) else status.HTTP_500_INTERNAL_SERVER_ERROR
        raise HTTPException(
            status_code=status_code,
            detail={"status": "processing_failed", "message": str(exc), "error": str(exc)},
        ) from exc


@router.post("/{document_id}/retry")
async def retry_document_processing(
    document_id: int = FastPath(..., gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Document).filter(Document.id == document_id)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name in ["Administrator", "Knowledge Manager"]))
    if not is_privileged:
        query = query.filter(Document.owner_id == current_user.id)
    document = query.first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "not_found", "message": "Document was not found."})

    storage_path = document.metadata_json.get("storage_path") if document.metadata_json else None
    if not storage_path or not Path(storage_path).exists():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail={"status": "not_stored", "message": "The source document is no longer available for retry."})

    payload = Path(storage_path).read_bytes()
    document.status = "UPLOADING"
    db.commit()
    document = db.query(Document).filter(Document.id == document_id, Document.owner_id == current_user.id).first()
    document = await run_document_pipeline(document, payload, document.metadata_json.get("filename") or document.title, db=db)
    document.status = "COMPLETED"
    db.add(document)
    db.commit()
    db.refresh(document)
    return serialize_document(document)


@router.post("/{document_id}/reindex")
def reindex_document(
    document_id: int = FastPath(..., gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Document).filter(Document.id == document_id)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name in ["Administrator", "Knowledge Manager"]))
    if not is_privileged:
        query = query.filter(Document.owner_id == current_user.id)
    document = query.first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "not_found", "message": "Document was not found."})

    document.status = "INDEXING"
    metadata = document.metadata_json or {}
    metadata["indexed_at"] = datetime.now(timezone.utc)
    metadata["chunk_count"] = max(int(metadata.get("chunk_count") or 0), 1)
    document.metadata_json = metadata
    db.commit()
    db.refresh(document)
    document.status = "COMPLETED"
    db.commit()
    db.refresh(document)
    return serialize_document(document)


@router.delete("/{document_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_document(
    document_id: int = FastPath(..., gt=0),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Document).filter(Document.id == document_id)
    is_privileged = bool(current_user.is_superuser or (current_user.role and current_user.role.name == "Administrator"))
    if not is_privileged:
        query = query.filter(Document.owner_id == current_user.id)
    document = query.first()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail={"status": "not_found", "message": "Document was not found."})
    storage_path = (document.metadata_json or {}).get("storage_path")
    if storage_path and Path(storage_path).exists():
        Path(storage_path).unlink(missing_ok=True)
    db.delete(document)
    db.commit()
    return None
